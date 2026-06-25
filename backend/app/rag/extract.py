import logging
from pathlib import Path

import pypdf
from docx import Document as DocxDocument

logger = logging.getLogger(__name__)


def extract_text(file_path: Path) -> str:
    """Extract text content from PDF, DOCX, or TXT files."""
    suffix = file_path.suffix.lower()

    if suffix == ".pdf":
        return _extract_pdf(file_path)
    if suffix == ".docx":
        return _extract_docx(file_path)
    if suffix in (".txt", ".md", ".csv"):
        return _extract_txt(file_path)

    logger.warning("Unsupported file type: %s", suffix)
    return ""


def _extract_pdf(file_path: Path) -> str:
    try:
        reader = pypdf.PdfReader(str(file_path))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = "\n".join(pages).strip()
        if not text:
            logger.warning("PDF '%s' extracted 0 chars — may be scanned/image PDF (no OCR)", file_path.name)
        return text
    except Exception as e:
        logger.error("PDF extraction failed for '%s': %s", file_path.name, e)
        return ""


def _extract_docx(file_path: Path) -> str:
    try:
        doc = DocxDocument(str(file_path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]

        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))

        if doc.sections:
            for section in doc.sections:
                for para in section.header.paragraphs:
                    if para.text.strip():
                        parts.append(para.text.strip())
                for para in section.footer.paragraphs:
                    if para.text.strip():
                        parts.append(para.text.strip())

        text = "\n".join(parts).strip()
        if not text:
            logger.warning("DOCX '%s' extracted 0 chars", file_path.name)
        return text
    except Exception as e:
        logger.error("DOCX extraction failed for '%s': %s", file_path.name, e)
        return ""


def _extract_txt(file_path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1258", "vni"):
        try:
            return file_path.read_text(encoding=encoding).strip()
        except (UnicodeDecodeError, UnicodeError):
            continue
    logger.warning("TXT '%s' — could not decode with any known encoding", file_path.name)
    return file_path.read_bytes().decode("utf-8", errors="replace").strip()
